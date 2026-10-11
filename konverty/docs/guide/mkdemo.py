# Демонстрационная база: вымышленные люди, только для скриншотов инструкции
import json, os, random, datetime, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment
rnd = random.Random(20261011)
# Запуск: python3 mkdemo.py путь/к/base.json — из настоящей базы берутся только улицы, индексы и районы
G = os.path.dirname(os.path.abspath(__file__))
os.makedirs(os.path.join(G, 'demo'), exist_ok=True)
real = json.load(open(sys.argv[1], encoding='utf-8'))
D = real['districts']
# улица/индекс/район берём из реальной базы (это общедоступные сведения), дом/квартира/люди — вымышленные
places = [(r[0], r[4], r[5], r[6], r[7]) for r in real['rows'] if r[7]]
FAM = ['СМИРНОВ','ИВАНОВ','КУЗНЕЦОВ','ПОПОВ','ВАСИЛЬЕВ','ПЕТРОВ','СОКОЛОВ','МИХАЙЛОВ','НОВИКОВ','ФЁДОРОВ','МОРОЗОВ','ВОЛКОВ','АЛЕКСЕЕВ','ЛЕБЕДЕВ','СЕМЁНОВ','ЕГОРОВ','ПАВЛОВ','КОЗЛОВ','СТЕПАНОВ','НИКОЛАЕВ','ОРЛОВ','АНДРЕЕВ','МАКАРОВ','НИКИТИН','ЗАХАРОВ','ЗАЙЦЕВ','СОЛОВЬЁВ','БОРИСОВ','ЯКОВЛЕВ','ГРИГОРЬЕВ','РОМАНОВ','ВОРОБЬЁВ','СЕРГЕЕВ','КУЗЬМИН','ФРОЛОВ','АЛЕКСАНДРОВ','ДМИТРИЕВ','КОРОЛЁВ','ГУСЕВ','КИСЕЛЁВ','ИЛЬИН','МАКСИМОВ','ПОЛЯКОВ','СОРОКИН','ВИНОГРАДОВ','КОВАЛЁВ','БЕЛОВ','МЕДВЕДЕВ','АНТОНОВ','ТАРАСОВ','ЖУКОВ','БАРАНОВ','ФИЛИППОВ','КОМАРОВ','ДАВЫДОВ','БЕЛЯЕВ','ГЕРАСИМОВ','БОГДАНОВ','ОСИПОВ','СИДОРОВ']
IM_F = ['АННА','МАРИЯ','ГАЛИНА','НИНА','ВАЛЕНТИНА','ЛЮДМИЛА','ТАМАРА','ЗИНАИДА','ЛИДИЯ','ВЕРА','НАДЕЖДА','ЕВГЕНИЯ','РАИСА','АНТОНИНА','КЛАВДИЯ','ТАТЬЯНА','ЕКАТЕРИНА','ЛЮБОВЬ','МАРГАРИТА','АЛЕКСАНДРА']
IM_M = ['ИВАН','ПЁТР','НИКОЛАЙ','ВЛАДИМИР','АЛЕКСАНДР','БОРИС','ВИКТОР','ЮРИЙ','ГЕННАДИЙ','ВАСИЛИЙ','МИХАИЛ','АНАТОЛИЙ','ЛЕОНИД','ВАЛЕНТИН','ЕВГЕНИЙ']
OT = ['ИВАНОВ','ПЕТРОВ','НИКОЛАЕВ','ВЛАДИМИРОВ','АЛЕКСАНДРОВ','БОРИСОВ','ВИКТОРОВ','ВАСИЛЬЕВ','МИХАЙЛОВ','ФЁДОРОВ','СЕРГЕЕВ','ПАВЛОВ','СТЕПАНОВ','ГРИГОРЬЕВ','АНДРЕЕВ']
def fem(s): return s + 'А'
norm = lambda s: str(s or '').upper().replace('Ё', 'Е').strip()
# вымышленный человек не должен совпасть по ФИО ни с кем из настоящей базы (сами ФИО из неё никуда не попадают)
REAL_FIO = {(norm(r[1]), norm(r[2]), norm(r[3])) for r in real['rows']}
rows, used, renamed = [], set(), 0
def add(di, fam, im, ot, idx, town, np, st, house, korp, apt, g):
    global renamed
    if (norm(fam), norm(im), norm(ot)) in REAL_FIO:
        suf = 'НА' if g == 'Ж' else 'ИЧ'
        ot = next(o + suf for o in OT if (norm(fam), norm(im), norm(o + suf)) not in REAL_FIO and (fam, im, o + suf) not in used)
        renamed += 1
    key = (fam, im, ot)
    if key in used: return False
    used.add(key); rows.append([di, fam, im, ot, idx, town, np, st, house, korp, apt, g]); return True
# Особые записи для сценариев инструкции
def place_with(street_sub):
    return next(p for p in places if street_sub in p[4])
p = place_with('УЛИЦА ЛЕНСКАЯ'); add(p[0], 'СМИРНОВА', 'ГАЛИНА', 'ПЕТРОВНА', p[1], p[2], p[3], p[4], '12', '1 ЛИТЕРА А', '45', 'Ж')
p = place_with('ПРОСПЕКТ ПРОСВЕЩЕНИЯ'); add(p[0], 'КУЗНЕЦОВА', 'АННА', 'ИВАНОВНА', p[1], p[2], p[3], p[4], '30', 'ЛИТЕРА А', '12', 'Ж'); add(p[0], 'КУЗНЕЦОВА', 'АНТОНИНА', 'ИВАНОВНА', p[1], p[2], p[3], p[4], '30', 'ЛИТЕРА А', '87', 'Ж')
p = place_with('ПРОСПЕКТ КОСМОНАВТОВ'); add(p[0], 'ОРЛОВ', 'НИКОЛАЙ', 'ВАСИЛЬЕВИЧ', p[1], p[2], p[3], p[4], '28', '2 ЛИТЕРА А', '16', 'М')
p = place_with('УЛИЦА ДЫБЕНКО'); add(p[0], 'ВОРОБЬЁВА', 'ЛИДИЯ', 'МИХАЙЛОВНА', p[1], p[2], p[3], p[4], '8', '1', '120', 'Ж')
while len(rows) < 600:
    di, idx, town, np, st = rnd.choice(places)
    g = 'Ж' if rnd.random() < 0.78 else 'М'
    fam = rnd.choice(FAM); im = rnd.choice(IM_F if g == 'Ж' else IM_M); ot = rnd.choice(OT)
    if g == 'Ж': fam, ot = fem(fam), ot + 'НА'
    else: ot = ot + 'ИЧ'
    add(di, fam, im, ot, idx, town, np, st, str(rnd.randint(1, 120)), rnd.choice(['ЛИТЕРА А', '1 ЛИТЕРА А', '2 ЛИТЕРА А', '']), str(rnd.randint(1, 400)), g)
rows.sort(key=lambda r: (r[1], r[2], r[3]))
json.dump({'districts': D, 'rows': [r[:11] for r in rows], 'source': 'демонстрационная база (вымышленные люди)'}, open(os.path.join(G, 'demo/demo-base.json'), 'w', encoding='utf-8'), ensure_ascii=False)
# Excel той же структуры, что и рабочий файл
hdr = ['Наименование региона','Фамилия','Имя','Отчество','Индекс','Наименование субъекта Российской Федерации','Район/город','Населенный пункт','Улица','Дом','Корпус','Квартира','Дата рождения','Пол','Льготная категория','Код льготной категории','Отметка о вручении ("+" "-")','Основание невручения (о, п, с)','Пояснения оснований невручения','Район Санкт-Петербурга']
wb = Workbook(); ws = wb.active; ws.title = 'г. Санкт-Петербург'
thin = Side(style='thin', color='000000'); med = Side(style='medium', color='000000')
ws.append(hdr)
for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = Border(left=thin, right=thin, top=med, bottom=med)
ws.row_dimensions[1].height = 60
for r in rows:
    di, fam, im, ot, idx, town, np, st, house, korp, apt, g = r
    bd = datetime.date(rnd.randint(1925, 1944), rnd.randint(1, 12), rnd.randint(1, 28))
    ws.append(['г. Санкт-Петербург', fam, im, ot, int(idx) if idx.isdigit() else idx, 'г. Санкт-Петербург', town, np, st, house, korp, apt, bd, g, 'Житель блокадного Ленинграда', 3, None, None, None, D[di]])
for row in ws.iter_rows(min_row=2):
    for c in row: c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    row[12].number_format = 'DD.MM.YYYY'
widths = [23,20,14,20,8,19,16,17,40,7,11,9,13,5,30,9,12,18,21,22]
for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
ws.auto_filter.ref = 'A1:T1'
wb.save(os.path.join(G, 'demo/demo-base.xlsx'))
assert not {(norm(r[1]), norm(r[2]), norm(r[3])) for r in rows} & REAL_FIO
print('записей', len(rows), '| отчество заменено, чтобы не совпасть с настоящей базой:', renamed)
